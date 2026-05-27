"""Unit tests for perdiem.engine.report (PD-REP-002)."""
from __future__ import annotations

import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from perdiem.engine.models import Claim, CrewResult, DayVerdict
from perdiem.engine.report import format_periods, write_exceptions, write_master


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _crew(
    staff_id: str = "1234567",
    name: str = "John Doe",
    email: str = "john@test.com",
    periods: list[tuple[date, date]] | None = None,
    days: int = 2,
    total_thb: int = 800,
    remark: str | None = None,
) -> CrewResult:
    return CrewResult(
        staff_id=staff_id,
        name=name,
        email=email,
        periods=periods or [(date(2026, 2, 3), date(2026, 2, 4))],
        days=days,
        total_thb=total_thb,
        remark=remark,
    )


def _claim(
    staff_id: str = "1234567",
    name: str = "John Doe",
    roster_links: list[str] | None = None,
) -> Claim:
    return Claim(
        source="POSTING_BASE",
        source_row_ref="Posting!A2",
        timestamp=None,
        email="john@test.com",
        staff_id=staff_id,
        name=name,
        position=None,
        base="DMK",
        claim_type="Posting",
        claim_month="FEBRUARY 2026",
        claimed_days=[3, 4],
        roster_links=roster_links or ["https://drive.google.com/file/d/abc123"],
    )


def _verdict(d: date, verdict: str = "INVALID", rule: str = "R5", reason: str = "test") -> DayVerdict:
    return DayVerdict(claimed_date=d, verdict=verdict, rule=rule, reason=reason)


def _load_wb(data: bytes):
    return load_workbook(io.BytesIO(data))


# ---------------------------------------------------------------------------
# format_periods
# ---------------------------------------------------------------------------


def test_format_single_day():
    periods = [(date(2026, 2, 3), date(2026, 2, 3))]
    assert format_periods(periods) == "3 February 2026"


def test_format_range():
    periods = [(date(2026, 2, 3), date(2026, 2, 4))]
    assert format_periods(periods) == "3 February 2026 - 4 February 2026"


def test_format_multi_range_newline_joined():
    periods = [
        (date(2026, 1, 26), date(2026, 1, 27)),
        (date(2026, 2, 10), date(2026, 2, 11)),
    ]
    result = format_periods(periods)
    assert "26 January 2026 - 27 January 2026" in result
    assert "10 February 2026 - 11 February 2026" in result
    assert "\n" in result


def test_format_empty_returns_empty_string():
    assert format_periods([]) == ""


# ---------------------------------------------------------------------------
# write_master
# ---------------------------------------------------------------------------


def test_write_master_returns_bytes():
    data = write_master([_crew()], "FEBRUARY 2026")
    assert isinstance(data, bytes)
    assert len(data) > 0


def test_write_master_sheet_named_after_month():
    data = write_master([_crew()], "FEBRUARY 2026")
    wb = _load_wb(data)
    assert "FEBRUARY 2026" in wb.sheetnames


def test_write_master_has_header_row():
    data = write_master([_crew()], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    # Row 2 is the header row (row 1 is banner)
    headers = [ws.cell(row=2, column=c).value for c in range(1, 10)]
    assert "Item" in headers
    assert "New ID No." in headers
    assert "Employee name" in headers
    assert "Period" in headers
    assert "Days" in headers
    assert "Total Perdiem (THB)" in headers


def test_write_master_crew_row_data():
    data = write_master([_crew()], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    # Data starts at row 3
    assert ws.cell(row=3, column=1).value == 1          # Item
    assert ws.cell(row=3, column=2).value == "1234567"  # New ID No.
    assert ws.cell(row=3, column=3).value == "John Doe" # Employee name
    assert ws.cell(row=3, column=5).value == 2          # Days
    assert ws.cell(row=3, column=6).value == 800        # Total Perdiem


def test_write_master_period_text_in_cell():
    data = write_master([_crew()], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    period_cell = ws.cell(row=3, column=4).value
    assert "February 2026" in (period_cell or "")


def test_write_master_multi_range_period():
    r = _crew(
        periods=[
            (date(2026, 1, 26), date(2026, 1, 27)),
            (date(2026, 2, 10), date(2026, 2, 11)),
        ],
        days=4,
        total_thb=1600,
    )
    data = write_master([r], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    period_cell = ws.cell(row=3, column=4).value or ""
    assert "January" in period_cell
    assert "February" in period_cell


def test_write_master_empty_results_still_has_headers():
    data = write_master([], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    assert ws.cell(row=2, column=1).value == "Item"


def test_write_master_remark_written():
    r = _crew(remark="ตกเบิกเดือนมกราคม")
    data = write_master([r], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    assert "ตกเบิก" in (ws.cell(row=3, column=9).value or "")


def test_write_master_item_numbers_sequential():
    crew = [_crew(staff_id=f"100000{i}", name=f"Crew {i}") for i in range(3)]
    data = write_master(crew, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    items = [ws.cell(row=r, column=1).value for r in range(3, 6)]
    assert items == [1, 2, 3]


def test_write_master_multiple_crew():
    crew = [
        _crew(staff_id="1111111", name="Alice"),
        _crew(staff_id="2222222", name="Bob", periods=[(date(2026, 2, 5), date(2026, 2, 6))]),
    ]
    data = write_master(crew, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb["FEBRUARY 2026"]
    assert ws.cell(row=3, column=3).value == "Alice"
    assert ws.cell(row=4, column=3).value == "Bob"


# ---------------------------------------------------------------------------
# write_exceptions
# ---------------------------------------------------------------------------


def test_write_exceptions_returns_bytes():
    data = write_exceptions([], "FEBRUARY 2026")
    assert isinstance(data, bytes)
    assert len(data) > 0


def test_write_exceptions_has_header_row():
    data = write_exceptions([], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    headers = [ws.cell(row=2, column=c).value for c in range(1, 11)]
    assert "Staff ID" in headers
    assert "Verdict" in headers
    assert "Rule" in headers
    assert "Reason" in headers


def test_write_exceptions_includes_invalid():
    d = date(2026, 2, 3)
    cv = [(_claim(), [_verdict(d, verdict="INVALID", rule="R5", reason="ID mismatch")])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    # Row 3 is first data row (1=banner, 2=header)
    assert ws.cell(row=3, column=7).value == "INVALID"
    assert ws.cell(row=3, column=8).value == "R5"
    assert "ID mismatch" in (ws.cell(row=3, column=9).value or "")


def test_write_exceptions_includes_needs_review():
    d = date(2026, 2, 3)
    cv = [(_claim(), [_verdict(d, verdict="NEEDS_REVIEW", rule="R3")])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    assert ws.cell(row=3, column=7).value == "NEEDS_REVIEW"


def test_write_exceptions_excludes_valid_verdicts():
    d = date(2026, 2, 3)
    cv = [(_claim(), [_verdict(d, verdict="VALID", rule="R8", reason=None)])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    # No data rows (row 3 should be empty)
    assert ws.cell(row=3, column=1).value is None


def test_write_exceptions_excludes_valid_backclaim():
    d = date(2026, 2, 3)
    cv = [(_claim(), [_verdict(d, verdict="VALID_BACKCLAIM", rule="R8", reason=None)])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    assert ws.cell(row=3, column=1).value is None


def test_write_exceptions_roster_link_included():
    d = date(2026, 2, 3)
    cv = [(_claim(roster_links=["https://drive.google.com/file/d/abc"]),
           [_verdict(d)])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    link_cell = ws.cell(row=3, column=10).value or ""
    assert "drive.google.com" in link_cell


def test_write_exceptions_source_ref_included():
    d = date(2026, 2, 3)
    cv = [(_claim(), [_verdict(d)])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    assert ws.cell(row=3, column=5).value == "Posting!A2"


def test_write_exceptions_item_numbers_sequential():
    d3, d4 = date(2026, 2, 3), date(2026, 2, 4)
    cv = [(_claim(), [_verdict(d3), _verdict(d4)])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    assert ws.cell(row=3, column=1).value == 1
    assert ws.cell(row=4, column=1).value == 2


def test_write_exceptions_empty_produces_headers_only():
    data = write_exceptions([], "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "Item"
    assert ws.cell(row=3, column=1).value is None


def test_write_exceptions_claimed_date_in_iso_format():
    d = date(2026, 2, 3)
    cv = [(_claim(), [_verdict(d)])]
    data = write_exceptions(cv, "FEBRUARY 2026")
    wb = _load_wb(data)
    ws = wb.active
    assert ws.cell(row=3, column=6).value == "2026-02-03"
